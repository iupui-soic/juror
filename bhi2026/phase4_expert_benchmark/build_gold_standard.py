"""Freeze the adjudicated 100-pair gold standard from the two raters.

Adjudication rule (deterministic, human-only — no LLM judgment enters the
gold standard, so it stays a clean independent reference for Tier-2):

  rating:
    - raters agree                  -> that value
    - differ by 1 ("lead-pick")     -> Rater 1 (HI faculty, lead rater)
    - differ by >=2 ("adjudicate"):
        * if exactly one rater's (rating, methodology) is rubric-inconsistent
          (rating>=3 but methodology="no") -> take the OTHER rater
        * otherwise -> rounded mean of the two ratings, ties broken toward R1
    - Rater 1 abstained (pair_025, pair_035) -> Rater 2 (only human score)

  methodology:
    - follow the rater whose rating was adopted (lead R1 for agree/lead-pick/
      mean cases; R2 for the inconsistency-tiebreak and abstain cases), then
      enforce rubric consistency: adjudicated rating>=3 cannot be "no" -> set
      "partial".

This simultaneously fixes the 4 HARD rubric violations (rating-3/method-no),
which all coincide with >=2-point disagreements where R2 rated low.

Outputs: gold_standard.csv (audit columns) and writes the Adjudication sheet
into gold_standard.xlsx.
Run: python3 build_gold_standard.py
"""
import csv
from openpyxl import load_workbook, Workbook
import compute_kappa_complete as base

# All rater values and pair identities come from the single canonical workbook.
PAIRS_XLSX = "benchmark_pairs.xlsx"
METHOD_C = "../phase3_method_c/ratings.csv"


def hard_viol(rating, method):
    return rating is not None and method is not None and rating >= 3 and method == "no"


def pair_summary():
    """pair_id -> dict(capstone_id, pmid, bucket, cosine) from the AUTHORITATIVE
    benchmark_pairs.xlsx (the pairs the raters actually scored)."""
    wb = load_workbook(PAIRS_XLSX, data_only=True)
    ws = wb["Pair Summary"]
    h = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(h, row))
        out[d["pair_id"]] = {
            "capstone_id": d["capstone_id"],
            "pmid": str(d["pmid"]),
            "bucket": d["_internal_bucket"],
            "cosine": d["cosine_sim"],
        }
    return out


def method_c():
    out = {}
    with open(METHOD_C) as f:
        for r in csv.DictReader(f):
            out[(r["capstone_id"], str(r["pmid"]))] = {
                "rating": r["thematic_alignment"],
                "method": r["shared_methodology"],
            }
    return out


def adjudicate(p, r1, r2):
    """Return (adj_rating, adj_method, rule)."""
    a, b = r1["rating"], r2["rating"]
    am, bm = r1["method"], r2["method"]

    if a is None and b is not None:
        return b, bm, "R1-abstained->R2"
    if b is None and a is not None:
        return a, am, "R2-abstained->R1"
    if a is None and b is None:
        return None, None, "both-blank"

    if a == b:
        # agree on rating; methodology lead-picks R1 (consistency enforced below)
        return a, am, "agree"

    if abs(a - b) == 1:
        return a, am, "lead-pick-R1(1pt)"

    # >=2 point disagreement
    v1, v2 = hard_viol(a, am), hard_viol(b, bm)
    if v1 and not v2:
        return b, bm, "adj:R1-rubric-inconsistent->R2"
    if v2 and not v1:
        return a, am, "adj:R2-rubric-inconsistent->R1"
    # rounded mean, ties toward R1
    mean = (a + b) / 2.0
    lo, hi = int(mean), int(mean) + 1
    if mean == int(mean):
        adj = int(mean)
    else:  # .5 tie -> nearer to R1
        adj = hi if abs(hi - a) < abs(lo - a) else lo
    adj_m = am  # lead methodology
    return adj, adj_m, "adj:rounded-mean(tie->R1)"


def main():
    r1 = base.rater1_by_pair()
    r2 = base.rater2_by_pair()
    ps = pair_summary()
    mc = method_c()
    ids = sorted(set(r1) & set(r2))

    rows = []
    for p in ids:
        adj_r, adj_m, rule = adjudicate(p, r1[p], r2[p])
        # enforce rubric consistency on adjudicated methodology
        fixed = ""
        if adj_r is not None and adj_r >= 3 and adj_m == "no":
            adj_m, fixed = "partial", "method no->partial (rating>=3)"
        info = ps.get(p, {})
        key = (info.get("capstone_id"), info.get("pmid"))
        mcr = mc.get(key, {})
        rows.append({
            "pair_id": p,
            "capstone_id": info.get("capstone_id"),
            "pmid": info.get("pmid"),
            "bucket": info.get("bucket"),
            "cosine_sim": info.get("cosine"),
            "r1_rating": r1[p]["rating"], "r1_method": r1[p]["method"],
            "r2_rating": r2[p]["rating"], "r2_method": r2[p]["method"],
            "adj_rating": adj_r, "adj_method": adj_m,
            "adjudication_rule": rule,
            "rubric_fix": fixed,
            "method_c_rating": mcr.get("rating"),
            "method_c_method": mcr.get("method"),
        })

    cols = ["pair_id", "capstone_id", "pmid", "bucket", "cosine_sim",
            "r1_rating", "r1_method", "r2_rating", "r2_method",
            "adj_rating", "adj_method", "adjudication_rule", "rubric_fix",
            "method_c_rating", "method_c_method"]
    with open("gold_standard.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Adjudication"
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    wb.save("gold_standard.xlsx")

    # ---- summary of what the rule did ----
    from collections import Counter
    print(f"Gold standard frozen: {len(rows)} pairs -> gold_standard.csv / .xlsx")
    print("\nAdjudication rule usage:")
    for rule, n in Counter(r["adjudication_rule"] for r in rows).most_common():
        print(f"  {rule:<34} {n}")
    print(f"\nRubric fixes applied: {sum(1 for r in rows if r['rubric_fix'])}")
    print("\nThe 8 >=2-point adjudications:")
    for r in rows:
        if r["adjudication_rule"].startswith("adj"):
            print(f"  {r['pair_id']}  R1={r['r1_rating']}/{r['r1_method']:<7} "
                  f"R2={r['r2_rating']}/{r['r2_method']:<7} -> "
                  f"adj={r['adj_rating']}/{r['adj_method']:<7} [{r['bucket']}]  ({r['adjudication_rule']})")
    miss = [r["pair_id"] for r in rows if r["method_c_rating"] is None]
    print(f"\nPairs without a Method C rating (excluded from Tier-2): {miss}")


if __name__ == "__main__":
    main()
