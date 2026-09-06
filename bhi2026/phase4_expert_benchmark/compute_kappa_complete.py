"""Inter-rater agreement from benchmark_pairs.xlsx.

Reads both raters and the pair metadata from the single canonical workbook,
matching columns BY NAME (robust to column position):

  sheet "Rater 1" / "Rater 2" : pair_id | rating (0-4) | methodology (yes/partial/no)
  sheet "Pair Summary"        : pair_id | _internal_bucket | ...

Agreement is computed pairwise-complete (a pair is used only when both raters
supplied that value). Run: python3 compute_kappa_complete.py
"""
from collections import Counter
from openpyxl import load_workbook

BENCH = "benchmark_pairs.xlsx"

RATING_COL = "rating (0-4)"
METHOD_COL = "methodology (yes/partial/no)"
ID_COL = "pair_id"
BUCKET_COL = "_internal_bucket"


# ---------------------------------------------------------------- readers
def read_xlsx_sheet(path, sheet):
    """Return list of dict rows keyed by the header names of `sheet`."""
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, values))
            for values in ws.iter_rows(min_row=2, values_only=True)]


def norm_rating(v):
    if v is None or v == "":
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def norm_method(v):
    if v is None:
        return None
    s = str(v).strip().lower()
    return s if s in ("yes", "partial", "no") else None


def _ratings_from_sheet(sheet):
    """pair_id -> {'rating':int|None, 'method':str|None}."""
    out = {}
    for r in read_xlsx_sheet(BENCH, sheet):
        pid = r.get(ID_COL)
        if not pid:
            continue
        out[pid] = {"rating": norm_rating(r.get(RATING_COL)),
                    "method": norm_method(r.get(METHOD_COL))}
    return out


def rater1_by_pair():
    return _ratings_from_sheet("Rater 1")


def rater2_by_pair():
    return _ratings_from_sheet("Rater 2")


def bucket_by_pair():
    out = {}
    for r in read_xlsx_sheet(BENCH, "Pair Summary"):
        pid = r.get(ID_COL)
        if pid:
            out[pid] = r.get(BUCKET_COL)
    return out


# ------------------------------------------------------------------- kappa
def cohen_kappa(pairs, labels=None, weights=None):
    """Weighted Cohen's kappa over an ordered label list.

    pairs   : list of (a, b) with no missing values
    labels  : ordered labels; inferred + sorted if None
    weights : None (unweighted), 'linear', or 'quadratic'
    """
    if not pairs:
        return float("nan"), 0
    if labels is None:
        labels = sorted({a for a, _ in pairs} | {b for _, b in pairs},
                        key=lambda x: (str(type(x)), x))
    idx = {l: i for i, l in enumerate(labels)}
    k = len(labels)
    n = len(pairs)
    cm = [[0] * k for _ in range(k)]
    for a, b in pairs:
        cm[idx[a]][idx[b]] += 1
    row = [sum(r) for r in cm]
    col = [sum(cm[i][j] for i in range(k)) for j in range(k)]
    if weights == "linear":
        w = [[abs(i - j) / max(1, k - 1) for j in range(k)] for i in range(k)]
    elif weights == "quadratic":
        w = [[((i - j) / max(1, k - 1)) ** 2 for j in range(k)] for i in range(k)]
    else:
        w = [[0 if i == j else 1 for j in range(k)] for i in range(k)]
    po = sum((1 - w[i][j]) * cm[i][j] for i in range(k) for j in range(k)) / n
    pe = sum((1 - w[i][j]) * (row[i] * col[j] / (n * n))
             for i in range(k) for j in range(k))
    if abs(1 - pe) < 1e-12:
        return 1.0, n
    return (po - pe) / (1 - pe), n


def paired(r1, r2, ids, field):
    out = []
    for pid in ids:
        a = r1.get(pid, {}).get(field)
        b = r2.get(pid, {}).get(field)
        if a is not None and b is not None:
            out.append((a, b, pid))
    return out


# -------------------------------------------------------------------- main
def main():
    r1 = rater1_by_pair()
    r2 = rater2_by_pair()
    buckets = bucket_by_pair()
    ids = sorted(set(r1) & set(r2))

    print(f"Rater 1 unique pairs: {len(r1)}   Rater 2 unique pairs: {len(r2)}")
    print(f"Shared pair_ids:      {len(ids)}")

    # ----- rating (0-4) -----
    rp = paired(r1, r2, ids, "rating")
    missing = [pid for pid in ids
               if r1[pid]["rating"] is None or r2[pid]["rating"] is None]
    a_b = [(a, b) for a, b, _ in rp]
    labels04 = [0, 1, 2, 3, 4]
    print("\n=== Rating (0-4) ===")
    print(f"Pairwise-complete n = {len(rp)}"
          + (f"  (dropped, missing a rating: {missing})" if missing else ""))
    ku, n = cohen_kappa(a_b, labels=labels04)
    kl, _ = cohen_kappa(a_b, labels=labels04, weights="linear")
    kq, _ = cohen_kappa(a_b, labels=labels04, weights="quadratic")
    print(f"  Unweighted Cohen kappa : {ku:.4f}")
    print(f"  Linear-weighted kappa  : {kl:.4f}")
    print(f"  Quadratic-weighted     : {kq:.4f}")
    exact = sum(1 for a, b in a_b if a == b) / len(a_b)
    within1 = sum(1 for a, b in a_b if abs(a - b) <= 1) / len(a_b)
    print(f"  Exact agreement        : {exact*100:.1f}%")
    print(f"  Within-1 agreement     : {within1*100:.1f}%")

    print("\n  Rating distribution:")
    print(f"    Rater 1: {dict(sorted(Counter(a for a, _ in a_b).items()))}")
    print(f"    Rater 2: {dict(sorted(Counter(b for _, b in a_b).items()))}")

    print("\n  Confusion matrix (rows = Rater 1, cols = Rater 2):")
    cm = [[0] * 5 for _ in range(5)]
    for a, b in a_b:
        cm[a][b] += 1
    print("          " + " ".join(f"{l:>4}" for l in labels04))
    for li, r in zip(labels04, cm):
        print(f"    R1={li}  " + " ".join(f"{v:>4}" for v in r))

    # ----- per-bucket rating -----
    print("\n  Per-bucket unweighted kappa:")
    for bk in ["aligned", "ambiguous", "misaligned"]:
        sub = [(a, b) for a, b, pid in rp if buckets.get(pid) == bk]
        kb, nb = cohen_kappa(sub, labels=labels04)
        print(f"    {bk:<11} n={nb:<3} kappa={kb:.4f}")

    # ----- methodology (yes/partial/no) -----
    mp = paired(r1, r2, ids, "method")
    m_ab = [(a, b) for a, b, _ in mp]
    m_labels = ["no", "partial", "yes"]
    print("\n=== Methodology overlap (yes/partial/no) ===")
    print(f"Pairwise-complete n = {len(mp)}")
    kmu, _ = cohen_kappa(m_ab, labels=m_labels)
    print(f"  Unweighted kappa       : {kmu:.4f}")
    mexact = sum(1 for a, b in m_ab if a == b) / len(m_ab) if m_ab else float("nan")
    print(f"  Exact agreement        : {mexact*100:.1f}%")


if __name__ == "__main__":
    main()
