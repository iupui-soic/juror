"""Compute Cohen's kappa between Rater 1 and Rater 2 on benchmark_pairs.xlsx.

Run: python3 compute_kappa.py [path/to/benchmark_pairs.xlsx]

Outputs the per-bucket and overall agreement, weighted and unweighted Cohen's
kappa, plus a confusion matrix.
"""
import sys
from collections import Counter
from openpyxl import load_workbook

PATH = sys.argv[1] if len(sys.argv) > 1 else "benchmark_pairs.xlsx"

def get_col(ws, col_letter):
    return [c.value for c in ws[col_letter][1:]]  # skip header

def cohen_kappa(a, b, weights=None):
    """Weighted Cohen's kappa over integer scale."""
    pairs = [(int(x), int(y)) for x, y in zip(a, b) if x is not None and y is not None]
    if not pairs:
        return float("nan"), 0
    labels = sorted({x for x, _ in pairs} | {y for _, y in pairs})
    k = len(labels)
    idx = {l: i for i, l in enumerate(labels)}
    n = len(pairs)
    cm = [[0]*k for _ in range(k)]
    for x, y in pairs:
        cm[idx[x]][idx[y]] += 1
    row_sums = [sum(row) for row in cm]
    col_sums = [sum(cm[i][j] for i in range(k)) for j in range(k)]
    if weights == "linear":
        w = [[abs(i - j) / max(1, k - 1) for j in range(k)] for i in range(k)]
    elif weights == "quadratic":
        w = [[((i - j) / max(1, k - 1))**2 for j in range(k)] for i in range(k)]
    else:
        w = [[0 if i == j else 1 for j in range(k)] for i in range(k)]
    po = sum((1 - w[i][j]) * cm[i][j] for i in range(k) for j in range(k)) / n
    pe = sum((1 - w[i][j]) * (row_sums[i] * col_sums[j] / (n * n))
             for i in range(k) for j in range(k))
    if abs(1 - pe) < 1e-12:
        return 1.0, n
    return (po - pe) / (1 - pe), n

def main():
    wb = load_workbook(PATH, data_only=True)
    r1_ratings = get_col(wb["Rater 1"], "F")
    r2_ratings = get_col(wb["Rater 2"], "F")
    r1_method = get_col(wb["Rater 1"], "G")
    r2_method = get_col(wb["Rater 2"], "G")
    pair_ids = get_col(wb["Rater 1"], "A")
    buckets = get_col(wb["Pair Summary"], "B")
    bucket_by_pair = dict(zip(get_col(wb["Pair Summary"], "A"), buckets))

    # Unweighted
    k, n = cohen_kappa(r1_ratings, r2_ratings)
    print(f"Unweighted Cohen kappa (n={n}): {k:.4f}")
    # Linear-weighted
    k, _ = cohen_kappa(r1_ratings, r2_ratings, weights="linear")
    print(f"Linear-weighted   Cohen kappa: {k:.4f}")
    # Quadratic-weighted
    k, _ = cohen_kappa(r1_ratings, r2_ratings, weights="quadratic")
    print(f"Quadratic-weighted Cohen kappa: {k:.4f}")

    # Methodology overlap kappa (3-class). Map yes/partial/no -> 2/1/0.
    _mmap = {"no": 0, "partial": 1, "yes": 2}
    def _mc(m):
        return _mmap.get(str(m).strip().lower()) if m is not None else None
    k_m, n_m = cohen_kappa([_mc(m) for m in r1_method], [_mc(m) for m in r2_method])
    print(f"\nMethodology overlap unweighted kappa (n={n_m}): {k_m:.4f}")

    # Per-bucket
    print("\nPer-bucket unweighted kappa:")
    for bucket in ["aligned", "ambiguous", "misaligned"]:
        a, b = [], []
        for pid, ra, rb in zip(pair_ids, r1_ratings, r2_ratings):
            if pid is None: continue
            if bucket_by_pair.get(pid) == bucket and ra is not None and rb is not None:
                a.append(ra); b.append(rb)
        kb, nb = cohen_kappa(a, b)
        print(f"  {bucket:<12} n={nb}  kappa={kb:.4f}")

    # Rating distributions
    print("\nRating distributions:")
    print(f"  Rater 1: {dict(sorted(Counter([r for r in r1_ratings if r is not None]).items()))}")
    print(f"  Rater 2: {dict(sorted(Counter([r for r in r2_ratings if r is not None]).items()))}")

    # Confusion matrix
    print("\nConfusion matrix (rows = Rater 1, cols = Rater 2):")
    labels = sorted({int(x) for x in r1_ratings if x is not None}
                    | {int(x) for x in r2_ratings if x is not None})
    cm = [[0]*len(labels) for _ in labels]
    for a, b in zip(r1_ratings, r2_ratings):
        if a is None or b is None: continue
        cm[labels.index(int(a))][labels.index(int(b))] += 1
    header = "        " + " ".join(f"{l:>4}" for l in labels)
    print(header)
    for li, row in zip(labels, cm):
        print(f"  R1={li}  " + " ".join(f"{v:>4}" for v in row))

if __name__ == "__main__":
    main()
