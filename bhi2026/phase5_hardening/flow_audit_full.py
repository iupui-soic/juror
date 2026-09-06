"""
Full analysis of the filled 33-flow audit workbooks.

Extends score_audit.py with everything the revision reports:
  (A) adjudicated two-rater standard, built with the same deterministic rule as
      phase4_expert_benchmark/build_gold_standard.py
  (B) the Table IV estimator panel for inter-rater and JUROR-vs-adjudicated
  (C) every baseline (Method C, cosine, BTM x4) scored against the human standard,
      with Steiger tests and flow- and cluster-bootstrap intervals
  (D) reliability, attainable ceiling, clustering design effect, power
  (E) justification quality and the invented-overlap rate

Writes aggregates only (no capstone text) to out/flow_audit.json.
The filled workbooks quote IRB-restricted capstone text and stay out of git.

Run:  python3 flow_audit_full.py <rater1>.xlsx <rater2>.xlsx
"""
import sys, csv, json, math
from pathlib import Path
import numpy as np
from scipy import stats
from scipy.stats import rankdata
import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT / "phase4_expert_benchmark"))
import agreement_panel as ap

Q, MLAB = 5, ["no", "partial", "yes"]
OUT = HERE / "out"; OUT.mkdir(exist_ok=True)


# ----------------------------------------------------------------- loading
def read_ratings(path):
    ws = openpyxl.load_workbook(path, data_only=True)["1_Rate_blind"]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    out = {}
    for r in range(2, ws.max_row + 1):
        fid = ws.cell(r, h["flow_id"]).value
        if not fid:
            continue
        v = ws.cell(r, h["your_rating_0to4"]).value
        m = (ws.cell(r, h["your_shared_methodology"]).value or "").strip().lower()
        out[fid] = dict(rating=None if v in (None, "") else int(v),
                        method=m if m in MLAB else None)
    return out


def read_faith(path):
    ws = openpyxl.load_workbook(path, data_only=True)["2_Audit_justification"]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    out = {}
    for r in range(2, ws.max_row + 1):
        fid = ws.cell(r, h["flow_id"]).value
        if not fid:
            continue
        g = lambda k: ws.cell(r, h[k]).value
        out[fid] = dict(cap=g("faith_capstone_desc_210"), jrn=g("faith_journal_desc_210"),
                        rel=g("faith_relationship_210"),
                        flag=(g("method_flag_correct_yn") or "").strip().lower(),
                        hall=(g("hallucinated_overlap_yn") or "").strip().lower())
    return out


def load_reference():
    juror, paper = {}, set()
    for r in csv.DictReader(open(ROOT / "phase3_method_d/flow_ratings.csv")):
        fid = f"C{int(r['capstone_topic'])}-J{int(r['journal_topic'])}"
        juror[fid] = dict(rating=int(r["thematic_alignment"]),
                          method=r["shared_methodology"].strip().lower())
        paper.add(fid)
    for extra in ("phase6_robustness/out/r9_topk_ratings.csv",
                  "phase6_robustness/out/r5_new_flow_ratings_claude.csv"):
        p = ROOT / extra
        if not p.exists():
            continue
        for r in csv.DictReader(open(p, encoding="utf-8")):
            if r.get("thematic_alignment") in ("", None):
                continue
            fid = f"C{int(r['capstone_topic'])}-J{int(r['journal_topic'])}"
            juror.setdefault(fid, dict(rating=int(float(r["thematic_alignment"])),
                                       method=r["shared_methodology"].strip().lower()))
    base = {}
    for r in csv.DictReader(open(ROOT / "phase6_robustness/out/r4_btm_flows.csv")):
        fid = f"C{int(r['capstone_topic'])}-J{int(r['journal_topic'])}"
        base[fid] = dict(cap=int(r["capstone_topic"]),
                         method_c=float(r["method_c"]), cosine=float(r["cosine"]),
                         btm_cap=float(r["btm_S_cap_native"]),
                         btm_jrn=float(r["btm_S_jrn_native"]),
                         btm_mean=float(r["btm_S_mean"]), btm_max=float(r["btm_S_max"]))
    return juror, paper, base


# ------------------------------------------------------- adjudication rule
def hard_viol(r, m):
    return r is not None and m is not None and r >= 3 and m == "no"


def adjudicate(a, am, b, bm):
    """Same deterministic, human-only rule as build_gold_standard.py."""
    if a is None and b is not None: return b, bm, "R1-abstained->R2"
    if b is None and a is not None: return a, am, "R2-abstained->R1"
    if a is None and b is None:     return None, None, "both-blank"
    if a == b:                      return a, am, "agree"
    if abs(a - b) == 1:             return a, am, "lead-pick-R1(1pt)"
    v1, v2 = hard_viol(a, am), hard_viol(b, bm)
    if v1 and not v2: return b, bm, "adj:R1-rubric-inconsistent->R2"
    if v2 and not v1: return a, am, "adj:R2-rubric-inconsistent->R1"
    mean = (a + b) / 2.0
    lo, hi = int(mean), int(mean) + 1
    adj = int(mean) if mean == int(mean) else (hi if abs(hi - a) < abs(lo - a) else lo)
    return adj, am, "adj:rounded-mean(tie->R1)"


# ---------------------------------------------------------------- helpers
def sp(a, b):
    ra, rb = rankdata(a), rankdata(b)
    if ra.std() == 0 or rb.std() == 0:
        return float("nan")
    return float(((ra - ra.mean()) * (rb - rb.mean())).mean() / (ra.std() * rb.std()))


def fisher_ci(r, n, a=0.05):
    if abs(r) >= 1 or n <= 3:
        return (float("nan"), float("nan"))
    z, se, zc = np.arctanh(r), 1 / math.sqrt(n - 3), stats.norm.ppf(1 - a / 2)
    return tuple(round(float(v), 3) for v in np.tanh([z - zc * se, z + zc * se]))


def chance_within1(a, b, q=Q):
    pa = np.bincount(np.asarray(a, int), minlength=q) / len(a)
    pb = np.bincount(np.asarray(b, int), minlength=q) / len(b)
    return float(sum(pa[i] * pb[j] for i in range(q) for j in range(q) if abs(i - j) <= 1))


def steiger(rjh, rbh, rjb, n):
    """Steiger (1980) z for two dependent correlations sharing one variable."""
    zj, zb = np.arctanh(rjh), np.arctanh(rbh)
    rm2 = (rjh ** 2 + rbh ** 2) / 2
    f = min((1 - rjb) / (2 * (1 - rm2)), 1.0)
    h = (1 - f * rm2) / (1 - rm2)
    return float((zj - zb) * math.sqrt((n - 3) / (2 * (1 - rjb) * h)))


def icc1k(cols):
    X = np.asarray(cols, float).T
    n, k = X.shape
    row, g = X.mean(1), X.mean()
    MSB = k * ((row - g) ** 2).sum() / (n - 1)
    MSW = ((X - row[:, None]) ** 2).sum() / (n * (k - 1))
    return float((MSB - MSW) / MSB)


def panel(a, b):
    a, b = list(map(int, a)), list(map(int, b))
    n = len(a)
    cm = ap.confusion(a, b, Q)
    Wid, Wlin, Wq = (ap.agree_weights(Q, k) for k in ("identity", "linear", "quadratic"))
    ia, _ = ap.icc_two_way(a, b)
    return dict(n=n,
                exact=float(np.trace(cm) / n),
                within1=float(np.mean([abs(x - y) <= 1 for x, y in zip(a, b)])),
                within1_chance=chance_within1(a, b),
                kappa=float(ap.cohen(cm, Wid)),
                kappa_linear=float(ap.cohen(cm, Wlin)),
                kappa_quadratic=float(ap.cohen(cm, Wq)),
                pabak=float(ap.pabak(cm)),
                ac1=float(ap.gwet(cm, Wid)),
                ac2_linear=float(ap.gwet(cm, Wlin)),
                ac2_quadratic=float(ap.gwet(cm, Wq)),
                krippendorff_ordinal=float(ap.krippendorff(list(zip(a, b)), "ordinal", range(Q))),
                icc_2_1=float(ia),
                pearson=float(stats.pearsonr(a, b)[0]),
                spearman=sp(a, b),
                spearman_ci=fisher_ci(sp(a, b), n),
                mean_a=float(np.mean(a)), mean_b=float(np.mean(b)))


def methodology(a, b):
    li = {v: i for i, v in enumerate(MLAB)}
    aa = [li[x] for x, y in zip(a, b) if x in li and y in li]
    bb = [li[y] for x, y in zip(a, b) if x in li and y in li]
    if len(aa) < 3:
        return None
    cm = ap.confusion(aa, bb, 3)
    return dict(n=len(aa), exact=float(np.trace(cm) / len(aa)),
                kappa=float(ap.cohen(cm, ap.agree_weights(3, "identity"))),
                ac1=float(ap.gwet(cm, ap.agree_weights(3, "identity"))))


# ------------------------------------------------------------------- main
def main(p1, p2):
    r1, r2 = read_ratings(p1), read_ratings(p2)
    f1, f2 = read_faith(p1), read_faith(p2)
    juror, PAPER, base = load_reference()

    fids = sorted(f for f in r1
                  if f in r2 and f in juror
                  and r1[f]["rating"] is not None and r2[f]["rating"] is not None)
    adj, rules = {}, {}
    for f in fids:
        ar, am, rule = adjudicate(r1[f]["rating"], r1[f]["method"],
                                  r2[f]["rating"], r2[f]["method"])
        if ar is not None and ar >= 3 and am == "no":
            am = "partial"
        adj[f] = dict(rating=ar, method=am); rules[f] = rule

    paper = [f for f in fids if f in PAPER]
    res = {"n_flows": len(fids), "n_paper_flows": len(paper),
           "adjudication_rules": {k: sum(1 for v in rules.values() if v == k)
                                  for k in sorted(set(rules.values()))},
           "adjudication_collapses_to_lead_rater":
               all(v in ("agree", "lead-pick-R1(1pt)") for v in rules.values())}

    for name, fs in (("paper_flows_24", paper), ("all_flows_33", fids)):
        blk = {
            "inter_rater": panel([r1[f]["rating"] for f in fs], [r2[f]["rating"] for f in fs]),
            "inter_rater_methodology": methodology([r1[f]["method"] for f in fs],
                                                   [r2[f]["method"] for f in fs]),
            "juror_vs_adjudicated": panel([juror[f]["rating"] for f in fs],
                                          [adj[f]["rating"] for f in fs]),
            "juror_vs_adjudicated_methodology": methodology([juror[f]["method"] for f in fs],
                                                            [adj[f]["method"] for f in fs]),
            "juror_vs_rater1": sp([juror[f]["rating"] for f in fs], [r1[f]["rating"] for f in fs]),
            "juror_vs_rater2": sp([juror[f]["rating"] for f in fs], [r2[f]["rating"] for f in fs]),
            "juror_vs_rater_mean": sp([juror[f]["rating"] for f in fs],
                                      [(r1[f]["rating"] + r2[f]["rating"]) / 2 for f in fs]),
        }
        rel = icc1k([[r1[f]["rating"] for f in fs], [r2[f]["rating"] for f in fs]])
        ceiling = math.sqrt(max(rel, 0) * 0.97)
        blk["reliability"] = dict(human_icc_1k=rel, juror_icc=0.97, ceiling=ceiling,
                                  spearman_disattenuated=blk["juror_vs_adjudicated"]["spearman"] / ceiling)
        H = np.array([(r1[f]["rating"] + r2[f]["rating"]) / 2 for f in fs])
        J = np.array([juror[f]["rating"] for f in fs], float)
        lo, hi = int((J < H).sum()), int((J > H).sum())
        blk["direction"] = dict(juror_below=lo, juror_above=hi, equal=int((J == H).sum()),
                                sign_test_p=float(stats.binomtest(hi, lo + hi, 0.5).pvalue),
                                mean_gap=float(np.mean(J - H)))
        res[name] = blk

    # ---- baselines vs the human standard, on the 24 paper flows
    fs = [f for f in paper if f in base]
    Hadj = np.array([adj[f]["rating"] for f in fs], float)
    Hmean = np.array([(r1[f]["rating"] + r2[f]["rating"]) / 2 for f in fs])
    CAP = np.array([base[f]["cap"] for f in fs])
    capidx = [np.where(CAP == c)[0] for c in np.unique(CAP)]
    J = np.array([juror[f]["rating"] for f in fs], float)
    preds = {"juror": J, **{k: np.array([base[f][k] for f in fs], float)
                            for k in ("method_c", "cosine", "btm_cap", "btm_jrn",
                                      "btm_mean", "btm_max")}}
    rng = np.random.default_rng(7)
    NB = 6000
    flow_bs = rng.integers(0, len(fs), size=(NB, len(fs)))
    clus_bs = [np.concatenate([capidx[c] for c in rng.choice(len(capidx), len(capidx), replace=True)])
               for _ in range(NB)]

    bl = {}
    for k, v in preds.items():
        e = {"vs_adjudicated": sp(v, Hadj), "vs_rater_mean": sp(v, Hmean),
             "vs_method_c": sp(v, preds["method_c"]),
             "ci_vs_adjudicated": fisher_ci(sp(v, Hadj), len(fs))}
        if k != "juror":
            rjb = sp(J, v)
            e["delta_rho_juror_minus_this"] = sp(J, Hadj) - sp(v, Hadj)
            e["steiger_z"] = steiger(sp(J, Hadj), sp(v, Hadj), rjb, len(fs))
            e["steiger_p"] = float(2 * (1 - stats.norm.cdf(abs(e["steiger_z"]))))
            d = np.array([x for i in flow_bs
                          if not math.isnan(x := sp(J[i], Hadj[i]) - sp(v[i], Hadj[i]))])
            dc = np.array([x for i in clus_bs
                           if not math.isnan(x := sp(J[i], Hadj[i]) - sp(v[i], Hadj[i]))])
            e["flow_bootstrap_ci"] = [float(np.percentile(d, 2.5)), float(np.percentile(d, 97.5))]
            e["cluster_bootstrap_ci"] = [float(np.percentile(dc, 2.5)), float(np.percentile(dc, 97.5))]
            e["n_zero_of_24"] = int((v == 0).sum())
        bl[k] = e
    res["baselines_vs_human"] = bl

    # ---- clustering and power
    grand, m = Hadj.mean(), 3
    MSB = m * sum((Hadj[i].mean() - grand) ** 2 for i in capidx) / (len(capidx) - 1)
    MSW = sum(((Hadj[i] - Hadj[i].mean()) ** 2).sum() for i in capidx) / (len(Hadj) - len(capidx))
    icc = (MSB - MSW) / (MSB + (m - 1) * MSW)
    deff = 1 + (m - 1) * icc
    scen = {}
    for nm, nf, nc in (("audited_24", 24, 8), ("audited_33", 33, 8),
                       ("full_depth_sweep_71", 71, 8), ("every_edge_848", 848, 8),
                       ("capstone_k20_top3", 60, 20), ("capstone_k40_top3", 120, 40),
                       ("four_corpus_pairs", 96, 32)):
        d = 1 + (nf / nc - 1) * icc
        scen[nm] = dict(flows=nf, clusters=nc, design_effect=round(d, 2),
                        n_effective=round(nf / d, 1))
    rho = sp(J, Hadj)
    power = {}
    for dt in (0.05, 0.10, 0.15, 0.20, 0.30):
        power[f"delta_{dt:.2f}"] = {
            f"n_{n}": round(float(stats.norm.cdf(abs(steiger(rho, rho - dt, 0.75, n)) - 1.96)
                                  + stats.norm.cdf(-abs(steiger(rho, rho - dt, 0.75, n)) - 1.96)), 3)
            for n in (24, 33, 71, 150, 300)}
    res["clustering_and_power"] = dict(
        icc_within_capstone_topic=float(icc), design_effect=float(deff),
        n_effective=float(len(fs) / deff), scenarios=scen, power_dependent_rho=power,
        note="cluster bootstrap on 8 clusters is itself high-variance; "
             "cluster-robust inference normally wants 30-50 clusters")

    # ---- justification quality
    jq = {}
    for axis, nm in (("cap", "capstone_description"), ("jrn", "journal_description"),
                     ("rel", "relationship")):
        a = np.array([f1[f][axis] for f in fids], float)
        b = np.array([f2[f][axis] for f in fids], float)
        jq[nm] = dict(rater1_mean=float(a.mean()), rater2_mean=float(b.mean()),
                      pooled_mean=float(((a + b) / 2).mean()),
                      fully_faithful_both=float(np.mean((a == 2) & (b == 2))),
                      either_below_2=float(np.mean((a < 2) | (b < 2))))
    for nm, key in (("methodology_flag_correct", "flag"), ("invented_overlap", "hall")):
        a = [f1[f][key] for f in fids]; b = [f2[f][key] for f in fids]
        jq[nm] = dict(rater1=sum(x == "yes" for x in a), rater2=sum(x == "yes" for x in b),
                      both=sum(x == "yes" and y == "yes" for x, y in zip(a, b)),
                      n=len(fids),
                      flows=[f for f, x, y in zip(fids, a, b) if x == "yes" or y == "yes"]
                      if key == "hall" else None)
    keep = [f for f in fids if f1[f]["hall"] != "yes" and f2[f]["hall"] != "yes"]
    jq["spearman_excluding_invented_overlap"] = sp([juror[f]["rating"] for f in keep],
                                                   [adj[f]["rating"] for f in keep])
    res["justification_quality"] = jq

    res["per_flow"] = {f: dict(rater1=r1[f]["rating"], rater2=r2[f]["rating"],
                               adjudicated=adj[f]["rating"], juror=juror[f]["rating"],
                               paper_flow=f in PAPER) for f in fids}

    (OUT / "flow_audit.json").write_text(json.dumps(res, indent=1))
    print(json.dumps({k: v for k, v in res.items() if k != "per_flow"}, indent=1))
    print(f"\nwrote {OUT / 'flow_audit.json'}")
    return res


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
