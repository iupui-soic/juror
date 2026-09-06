"""Extract text from every .pdf / .pptx under OneDrive_extracted/ into a new folder."""
import json
import os
import sys
from pathlib import Path

import fitz  # PyMuPDF
from pptx import Presentation

SRC = Path("raw_data/OneDrive_extracted/Practicum and capstone work")
OUT_DIR = Path("raw_data/OneDrive_extracted_text")
TXT_DIR = OUT_DIR / "texts"


def extract_pdf(path: Path) -> str:
    parts = []
    with fitz.open(path) as doc:
        for page in doc:
            parts.append(page.get_text("text"))
    return "\n".join(parts).strip()


def extract_pptx(path: Path) -> str:
    parts = []
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        slide_lines = [f"--- Slide {i} ---"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in p.runs).strip()
                    if line:
                        slide_lines.append(line)
            if shape.shape_type == 19 and getattr(shape, "has_table", False):  # table
                for row in shape.table.rows:
                    for cell in row.cells:
                        t = cell.text.strip()
                        if t:
                            slide_lines.append(t)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            slide_lines.append("[Notes] " + slide.notes_slide.notes_text_frame.text.strip())
        parts.append("\n".join(slide_lines))
    return "\n".join(parts).strip()


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    TXT_DIR.mkdir(parents=True, exist_ok=True)

    files = sorted(
        p for p in SRC.rglob("*")
        if p.is_file() and p.suffix.lower() in {".pdf", ".pptx"}
    )
    print(f"Found {len(files)} files")

    results = {}
    errors = []
    for i, path in enumerate(files, 1):
        rel = path.relative_to(SRC).as_posix()
        try:
            if path.suffix.lower() == ".pdf":
                text = extract_pdf(path)
            else:
                text = extract_pptx(path)
        except Exception as e:
            errors.append({"file": rel, "error": f"{type(e).__name__}: {e}"})
            text = ""
        results[rel] = text

        # write one .txt mirroring the folder layout
        out_txt = TXT_DIR / (rel + ".txt")
        out_txt.parent.mkdir(parents=True, exist_ok=True)
        out_txt.write_text(text, encoding="utf-8")

        if i % 25 == 0 or i == len(files):
            print(f"  {i}/{len(files)}  ({rel})")

    # consolidated json
    (OUT_DIR / "extracted_text.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    # errors log
    if errors:
        (OUT_DIR / "errors.json").write_text(
            json.dumps(errors, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    print(f"\nWrote {len(results)} text files to {TXT_DIR}")
    print(f"Wrote consolidated JSON to {OUT_DIR / 'extracted_text.json'}")
    if errors:
        print(f"{len(errors)} file(s) errored — see errors.json")


if __name__ == "__main__":
    main()
